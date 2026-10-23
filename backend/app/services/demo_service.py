import os
import shutil
import logging
from pathlib import Path

from backend.app.config import DEMO_DIR

logger = logging.getLogger("sloth.demo")

SAMPLE_FILES = [
    {
        "filename": "DSA_lecture4_graphs_and_trees_v2_final_FINAL.txt",
        "content": """Data Structures and Algorithms - Lecture 4: Graph Traversals & Trees
Instructor: Dr. Alan Turing
Date: October 14, 2026

1. Graph Representations:
- Adjacency Matrix: O(V^2) space, O(1) edge lookup. Best for dense graphs.
- Adjacency List: O(V + E) space, efficient iteration over neighbors. Best for sparse graphs.

2. Breadth-First Search (BFS):
- Uses a Queue (FIFO).
- Finds the shortest path in unweighted graphs.
- Time Complexity: O(V + E).
- Space Complexity: O(V) for visited set and queue.

3. Depth-First Search (DFS):
- Uses recursion or a Stack (LIFO).
- Useful for cycle detection, topological sorting, and finding connected components.
- Time Complexity: O(V + E).

4. Action Items & Assignments:
- Complete Homework 2 on Minimum Spanning Trees (Kruskal and Prim).
- Submission deadline: October 29, 2026 at 11:59 PM.
"""
    },
    {
        "filename": "arxiv_2310_transformer_attention_mechanisms_survey.txt",
        "content": """A Survey on Efficient Self-Attention in Large Language Models
Authors: S. Vaswani, J. Devlin, et al.
Abstract:
Standard transformer models rely on scaled dot-product self-attention with O(N^2) complexity in sequence length. In this survey, we categorize linear attention, FlashAttention, and sparse attention approaches. We highlight hardware-aware optimizations that achieve 3x throughput improvements on modern GPUs.

Key Takeaways:
1. Standard attention: Attention(Q, K, V) = softmax(QK^T / sqrt(d_k)) V.
2. Memory bandwidth is typically the primary bottleneck during autoregressive decoding.
3. FlashAttention reduces HBM memory read/writes by tiling computation in SRAM.

Action Items:
- Benchmark kernel throughput on Qwen 2.5 and Llama architectures.
- Review token throughput metrics by November 12, 2026.
"""
    },
    {
        "filename": "october_2026_office_supplies_receipts_tax.csv",
        "content": """Date,Vendor,Category,Amount_USD,Tax_Deductible,Notes
2026-10-02,AWS Cloud Hosting,Infrastructure,142.50,Yes,Production server compute
2026-10-05,Office Depot,Supplies,38.99,Yes,Notebooks and dry erase markers
2026-10-10,Stripe Terminal,Payment Processing,15.20,Yes,Merchant transaction fees
2026-10-15,Peet's Coffee,Meals & Entertainment,24.50,No,Team sprint sync coffee
2026-10-21,GitHub Enterprise,Software Subscriptions,84.00,Yes,Team licenses for Hacktoberfest
"""
    },
    {
        "filename": "sloth_agent_runtime_config_v1_debug.json",
        "content": """{
  "system": "SLOTH Local AI Agent",
  "version": "1.0.0",
  "local_first": true,
  "inference_backend": "Ollama",
  "default_model": "qwen2.5:3b",
  "settings": {
    "concurrency_limit": 2,
    "temperature": 0.2,
    "keep_alive": "5m"
  },
  "modules": ["scanner", "organizer", "digest", "relax_mode"]
}"""
    },
    {
        "filename": "random_download_838291_messy_name.txt",
        "content": """Personal Fitness & Marathon Training Log - Week 4
Target: Sub-4 hour marathon in April.
- Monday: 8 km tempo run at 5:10/km pace. Felt strong.
- Wednesday: Interval training: 6 x 800m with 90s recovery.
- Friday: Rest and core strength exercises.
- Sunday: Long run 22 km. Hydration strategy: electrolytes every 5 km.

Note to self: Pick up new running shoes before the 30 km long run on November 8, 2026.
"""
    }
]

# Minimal valid 1x1 PNG bytes for testing media handling
TINY_PNG_BYTES = bytes([
    0x89, 0x50, 0x4E, 0x47, 0x0D, 0x0A, 0x1A, 0x0A, 0x00, 0x00, 0x00, 0x0D,
    0x49, 0x48, 0x44, 0x52, 0x00, 0x00, 0x00, 0x01, 0x00, 0x00, 0x00, 0x01,
    0x08, 0x06, 0x00, 0x00, 0x00, 0x1F, 0x15, 0xC4, 0x89, 0x00, 0x00, 0x00,
    0x0A, 0x49, 0x44, 0x41, 0x54, 0x78, 0x9C, 0x63, 0x00, 0x01, 0x00, 0x00,
    0x05, 0x00, 0x01, 0x0D, 0x0A, 0x2D, 0xB4, 0x00, 0x00, 0x00, 0x00, 0x49,
    0x45, 0x4E, 0x44, 0xAE, 0x42, 0x60, 0x82
])

def create_sample_pdf(dest_path: Path):
    """Create a sample PDF using reportlab or fallback."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(dest_path), pagesize=letter)
        c.drawString(72, 750, "CS301: Advanced Operating Systems & Concurrency")
        c.drawString(72, 730, "Lecture Notes - Deadlock Detection & Banker's Algorithm")
        c.drawString(72, 700, "1. Necessary Conditions for Coffman Deadlocks:")
        c.drawString(90, 680, "- Mutual Exclusion")
        c.drawString(90, 665, "- Hold and Wait")
        c.drawString(90, 650, "- No Preemption")
        c.drawString(90, 635, "- Circular Wait")
        c.drawString(72, 605, "2. Key Deadlines:")
        c.drawString(90, 585, "- Lab 3 Submission due November 15, 2026")
        c.save()
    except Exception as e:
        logger.warning(f"Could not use reportlab for demo PDF: {e}")
        # Write text with pdf extension if reportlab fails
        with open(dest_path, "w", encoding="utf-8") as f:
            f.write("CS301 Operating Systems Notes - Deadlock Detection.")

def create_sample_docx(dest_path: Path):
    """Create a sample DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("Q4 2026 Product Launch & Sprint Brief", 0)
        doc.add_paragraph("SLOTH AI Agent - Autonomous Desktop Productivity Tool.")
        doc.add_heading("Key Deliverables", level=1)
        doc.add_paragraph("1. Smart Folder Organizer with 2-phase verification.")
        doc.add_paragraph("2. Document Digest with executive summaries and action items.")
        doc.add_paragraph("3. Relax Mode for smooth background batch execution.")
        doc.add_heading("Milestones and Deadlines", level=1)
        doc.add_paragraph("Internal Beta Launch: October 28, 2026.")
        doc.add_paragraph("Hacktoberfest Submission Window: October 31, 2026.")
        doc.save(str(dest_path))
    except Exception as e:
        logger.warning(f"Could not use docx for demo DOCX: {e}")

def init_demo_sandbox(force_reset: bool = True) -> str:
    """Create or reset sample files in the demo sandbox directory."""
    demo_path = DEMO_DIR.resolve()
    
    if force_reset and demo_path.exists():
        shutil.rmtree(demo_path)
        
    demo_path.mkdir(parents=True, exist_ok=True)
    
    # Write text and data sample files
    for item in SAMPLE_FILES:
        file_path = demo_path / item["filename"]
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(item["content"])
            
    # Write sample media
    png_path = demo_path / "screenshot_20261020_raw_unnamed.png"
    with open(png_path, "wb") as f:
        f.write(TINY_PNG_BYTES)
        
    # Write sample PDF
    pdf_path = demo_path / "CS301_OS_lecture_notes_deadlocks_v3.pdf"
    create_sample_pdf(pdf_path)
    
    # Write sample DOCX
    docx_path = demo_path / "product_launch_brief_q4_confidential.docx"
    create_sample_docx(docx_path)
    
    logger.info("Demo sandbox initialized with 8 sample files at: %s", demo_path)
    return str(demo_path)
